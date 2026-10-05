"""
Unit and integration tests for app/core/doc_parser.py.
Tests PDF section heading detection, TOC suppression, running footer removal,
and DOCX heading style preservation as Markdown.
"""

import sys
from pathlib import Path
import pytest

# Add backend to sys.path
backend_dir = Path(__file__).resolve().parent.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from app.core.doc_parser import parse_document, _detect_pdf_headings_and_convert_to_md
from app.core.doc_chunker import chunk_document


def test_pdf_numbered_headings_detected():
    """Test 1: Multi-page PDF text with 2+ H1 headings returns 'heading_aware'."""
    page1 = "1. Executive Summary\nThis is the executive summary section."
    page2 = "2. Product Overview\nThis section covers product overview.\n2.1 Architecture\nArchitecture details."
    
    text, method = _detect_pdf_headings_and_convert_to_md([page1, page2])
    assert method == "heading_aware"
    assert "# 1. Executive Summary" in text
    assert "# 2. Product Overview" in text
    assert "## 2.1 Architecture" in text


def test_pdf_no_headings_falls_back():
    """Test 2: Plain text PDF without numbered section structure falls back to 'fixed_window'."""
    page1 = "This is a unstructured plain document without any numbered headings.\nJust sentences flowing across paragraphs."
    page2 = "Second page continues the prose without any clear section headers."
    
    text, method = _detect_pdf_headings_and_convert_to_md([page1, page2])
    assert method == "fixed_window"
    assert "# " not in text


def test_pdf_insufficient_headings_falls_back():
    """Test 3: PDF text with only 1 H1 heading falls back to 'fixed_window' (fallback gate)."""
    page1 = "1. Single Section Header\nAll the text of this brief document is under section 1."
    page2 = "More text continuing under section 1 without a second top level heading."
    
    text, method = _detect_pdf_headings_and_convert_to_md([page1, page2])
    assert method == "fixed_window"


def test_pdf_numbered_h1_h2_h3_levels():
    """Test 4: Correctly converts 1. -> H1 (#), 1.1 -> H2 (##), 1.1.1 -> H3 (###)."""
    page1 = "1. First Top Level Heading\nContent for section 1."
    page2 = "2. Second Top Level Heading\nContent 2.\n2.1 Subsection\nSubcontent.\n2.1.1 Deep Detail\nDetail text."
    
    text, method = _detect_pdf_headings_and_convert_to_md([page1, page2])
    assert method == "heading_aware"
    assert "# 1. First Top Level Heading" in text
    assert "# 2. Second Top Level Heading" in text
    assert "## 2.1 Subsection" in text
    assert "### 2.1.1 Deep Detail" in text


def test_pdf_footer_lines_stripped():
    """Test 5: Running footer lines appearing on >= 2 pages are stripped from all output."""
    footer = "ChatDoc PRD v2.4 — Confidential Page"
    page1 = f"1. Introduction\nPage 1 text.\n{footer}"
    page2 = f"2. Architecture\nPage 2 text.\n{footer}"
    
    text, method = _detect_pdf_headings_and_convert_to_md([page1, page2])
    assert footer not in text


def test_pdf_toc_heading_interaction():
    """
    Test 6 (Key Integration Test - Decision 3):
    Simulates realistic multi-page PDF structure with identical section numbers in TOC and body.
    Verifies TOC entries are suppressed, body headings produce correct heading_path,
    and running footers are stripped.
    """
    footer = "ChatDoc PRD v2.4 — Confidential"
    
    # Page 1: Title + TOC entries ending in bare page numbers
    page1 = f"""
PRODUCT REQUIREMENT DOCUMENT
ChatDoc Enterprise RAG
2. Table of Contents
1. Executive Summary & Overview 2
5. User Personas & Key Workflows 4
5.1 Primary User Personas 4
6. Functional Requirements 5
{footer}
""".strip()

    # Page 2: Body with real headings (no page number suffix)
    page2 = f"""
1. Executive Summary & Overview
ChatDoc transforms knowledge into queryable intelligence.

5. User Personas & Key Workflows
5.1 Primary User Personas
Persona A: Legal & Compliance Analyst
Needs zero-hallucination answers.
{footer}
""".strip()

    # Page 3: Body continued with next H1
    page3 = f"""
6. Functional Requirements
FR-101 Document Processing
FR-102 Text Chunking
{footer}
""".strip()

    converted_text, method = _detect_pdf_headings_and_convert_to_md([page1, page2, page3])
    assert method == "heading_aware"

    # Run converted text through chunk_document
    chunks = chunk_document(
        text=converted_text,
        doc_id="bdoc:test/ChatDoc_PRD.pdf",
        filename="ChatDoc_PRD.pdf",
        chunking_method="heading_aware"
    )

    # 1. Verify running footer is absent from all chunk content
    for c in chunks:
        assert "Confidential" not in c.text

    # 2. Extract all heading_paths across chunks
    all_heading_paths = [c.heading_path for c in chunks]

    # 3. Assert real body headings produced correct heading_path structures
    assert ["5. User Personas & Key Workflows"] in all_heading_paths or \
           any(p == ["5. User Personas & Key Workflows", "5.1 Primary User Personas"] for p in all_heading_paths)
    
    assert ["6. Functional Requirements"] in all_heading_paths or \
           any("6. Functional Requirements" in p for p in all_heading_paths)

    # 4. Assert TOC lines with trailing page numbers did not create spurious standalone chunks
    for c in chunks:
        assert not (len(c.heading_path) == 1 and c.heading_path[0].endswith(" 4"))


def test_docx_parse_returns_heading_aware_method(tmp_path):
    """Test 7: parse_document returns heading_aware for .docx extension when docx is installed."""
    docx_file = tmp_path / "test_doc.docx"
    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("Dummy content")
        doc.save(str(docx_file))
        _, method = parse_document(docx_file, "test_doc.docx")
        assert method == "heading_aware"
    except ImportError:
        docx_file.write_text("Dummy content")
        _, method = parse_document(docx_file, "test_doc.docx")
        assert method == "fixed_window"


def test_docx_heading_styles_preserved_as_markdown(tmp_path):
    """
    Test 8 (Milestone 21):
    Verifies that parse_document for .docx converts paragraph heading styles
    (Heading 1, Heading 2, Heading 3+) into Markdown '#', '##', '###' prefixes.
    """
    docx_file = tmp_path / "sample.docx"
    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("Heading 1 Text", style="Heading 1")
        doc.add_paragraph("Body content for section 1")
        doc.add_paragraph("Heading 2 Subtitle", style="Heading 2")
        doc.add_paragraph("Body content for section 2")
        doc.save(str(docx_file))

        text, method = parse_document(docx_file, "sample.docx")
        assert method == "heading_aware"
        assert "# Heading 1 Text" in text
        assert "## Heading 2 Subtitle" in text
    except ImportError:
        pytest.skip("python-docx not installed")


def test_docx_heading_aware_full_pipeline(tmp_path):
    """
    Test 9 (Milestone 21):
    Full pipeline test: parse_document on DOCX -> chunk_document produces
    heading_aware chunks with 3-tier heading_path metadata, section levels, and parent links.
    """
    docx_file = tmp_path / "full_doc.docx"
    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("AutoKT Architecture Specification", style="Heading 1")
        doc.add_paragraph("AutoKT is an automated knowledge transfer system designed to ingest and index code repositories.")
        doc.add_paragraph("System Overview", style="Heading 2")
        doc.add_paragraph("The backend is built with FastAPI, Neo4j graph database, and ChromaDB vector stores.")
        doc.add_paragraph("Document Ingestion Subsystem", style="Heading 2")
        doc.add_paragraph("Documents uploaded via API are parsed into Markdown format and chunked using section structure.")
        doc.add_paragraph("Heading-Aware DOCX Parsing", style="Heading 3")
        doc.add_paragraph("As of Milestone 21, DOCX paragraph styles are preserved as Markdown headers (# / ## / ###).")
        doc.save(str(docx_file))

        text, method = parse_document(docx_file, "full_doc.docx")
        assert method == "heading_aware"

        chunks = chunk_document(
            text=text,
            doc_id="doc:docx_test",
            filename="full_doc.docx",
            chunking_method=method,
        )

        assert len(chunks) == 4
        assert all(c.chunking_method == "heading_aware" for c in chunks)

        h1_chunk = next(c for c in chunks if c.heading_path == ["AutoKT Architecture Specification"])
        assert h1_chunk.section_level == 1
        assert h1_chunk.parent_section_id == ""

        h2_sys = next(c for c in chunks if c.heading_path == ["AutoKT Architecture Specification", "System Overview"])
        assert h2_sys.section_level == 2
        assert h2_sys.parent_section_id == h1_chunk.chunk_id

        h2_ing = next(c for c in chunks if c.heading_path == ["AutoKT Architecture Specification", "Document Ingestion Subsystem"])
        assert h2_ing.section_level == 2
        assert h2_ing.parent_section_id == h1_chunk.chunk_id

        h3_docx = next(c for c in chunks if c.heading_path == ["AutoKT Architecture Specification", "Document Ingestion Subsystem", "Heading-Aware DOCX Parsing"])
        assert h3_docx.section_level == 3
        assert h3_docx.parent_section_id == h2_ing.chunk_id
    except ImportError:
        pytest.skip("python-docx not installed")


def test_docx_no_headings_falls_back_to_fixed_window(tmp_path):
    """
    Test 10 (Milestone 21):
    Verifies that a DOCX file with only plain body text (no heading styles)
    falls back to fixed_window chunking inside chunk_document.
    """
    docx_file = tmp_path / "plain_doc.docx"
    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("This is plain body text paragraph 1 without heading styles.")
        doc.add_paragraph("This is plain body text paragraph 2 without heading styles.")
        doc.save(str(docx_file))

        text, method = parse_document(docx_file, "plain_doc.docx")
        assert method == "heading_aware"

        chunks = chunk_document(
            text=text,
            doc_id="doc:plain_docx",
            filename="plain_doc.docx",
            chunking_method=method,
        )

        assert len(chunks) >= 1
        assert all(c.chunking_method == "fixed_window" for c in chunks)
    except ImportError:
        pytest.skip("python-docx not installed")


